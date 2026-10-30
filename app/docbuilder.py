"""ชุดเครื่องมือสร้างเอกสาร Word รูปแบบงานวิจัยภาษาไทย (TH SarabunPSK 16 pt)
และแปลงเป็น PDF ผ่าน Microsoft Word

บล็อกเนื้อหา (block) ที่รองรับ:
  ("chapter", "บทที่ 1", "บทนำ")      หัวบท จัดกึ่งกลาง
  ("title", "ข้อความ")                  ข้อความจัดกึ่งกลางตัวหนา
  ("h1", "1.1 หัวข้อ")                  หัวข้อหลัก
  ("h2", "1.1.1 หัวข้อย่อย")            หัวข้อย่อย
  ("hd", "1. หลักการและเหตุผล")         หัวข้อของแบบฟอร์ม (ตัวหนา ขนาดเท่าเนื้อหา เว้นระยะก่อนหัวข้อ)
  ("p", "ย่อหน้า")                      ย่อหน้าเยื้องบรรทัดแรก
  ("pn", "ย่อหน้า")                     ย่อหน้าไม่เยื้อง
  ("item", "1) ข้อความ", level)         รายการแบบเยื้อง
  ("table", {caption, headers, rows, widths, note})
  ("gantt", {months, weeks, rows})      ตารางขั้นตอนการดำเนินงานรายสัปดาห์ (แรเงาช่วงที่ทำ)
  ("sign", {rows})                      ช่องลงนามแบบ 2 คอลัมน์ ไม่มีเส้นตาราง
  ("fill", n)                           เส้นประเว้นให้เขียนเพิ่ม n บรรทัด
  ("ref", "รายการอ้างอิง")              บรรณานุกรมแบบเยื้องแขวน
  ("pb",)                               ขึ้นหน้าใหม่
  ("blank",)                            บรรทัดว่าง

การจัดรูปแบบในข้อความ:  *ตัวเอียง*   **ตัวหนา**   ⟦ช่องเติมข้อมูล⟧ (ไฮไลต์สีเหลือง)
"""
from __future__ import annotations

import re
import threading
from functools import lru_cache
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "TH SarabunPSK"
BASE_SIZE = 16
TEXT_WIDTH = Inches(8.27 - 1.5 - 1.0)  # A4 - left - right

TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|⟦.+?⟧)")


# ---------------------------------------------------------------- font utils
def _set_run_font(run, size=BASE_SIZE, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs")
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(size * 2)))
    if bold is not None:
        run.bold = bold
        b = OxmlElement("w:bCs")
        if bold:
            rpr.append(b)
    if italic is not None:
        run.italic = italic
        if italic:
            rpr.append(OxmlElement("w:iCs"))


def _thai_justify(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    jc.set(qn("w:val"), "thaiDistribute")


def add_rich_text(paragraph, text, size=BASE_SIZE, bold=False):
    """เพิ่มข้อความที่มี *ตัวเอียง* **ตัวหนา** และ ⟦ช่องเติม⟧"""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(thai_wrap(part[2:-2]))
            _set_run_font(r, size, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(thai_wrap(part[1:-1]))
            _set_run_font(r, size, bold=bold, italic=True)
        elif part.startswith("⟦"):
            r = paragraph.add_run("[" + thai_wrap(part[1:-1]) + "]")
            _set_run_font(r, size, bold=bold)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = paragraph.add_run(thai_wrap(part))
            _set_run_font(r, size, bold=bold)


ZWSP = "\u200b"
_THAI = re.compile(r"[\u0E00-\u0E7F]")


_extra_words: set[str] = set()


def add_words(words):
    """เพิ่มคำ (เช่น ชื่อผู้แต่ง) ไม่ให้ถูกตัดกลางคำ"""
    new = {w for w in words if w and w not in _extra_words}
    if new:
        _extra_words.update(new)
        _thai_dict.cache_clear()
        thai_wrap.cache_clear()


@lru_cache(maxsize=1)
def _thai_dict():
    from pythainlp.corpus.common import thai_words
    from pythainlp.util import Trie

    return Trie(set(thai_words()) | _extra_words)


@lru_cache(maxsize=4096)
def thai_wrap(text: str) -> str:
    """แทรก zero-width space ระหว่างคำภาษาไทย เพื่อให้ Word ตัดบรรทัดได้ถูกตำแหน่ง"""
    if not _THAI.search(text):
        return text
    try:
        from pythainlp.tokenize import word_tokenize
    except ImportError:  # ไม่มี pythainlp -> คืนข้อความเดิม
        return text
    out = []
    prev = ""
    for tok in word_tokenize(text, engine="newmm", keep_whitespace=True, custom_dict=_thai_dict()):
        if out and _THAI.match(tok[:1] or " ") and _THAI.match(prev[-1:] or " "):
            out.append(ZWSP)
        out.append(tok)
        prev = tok
    return "".join(out)


def _fmt(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


# ---------------------------------------------------------------- document
def new_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = Inches(1.5)
    sec.left_margin = Inches(1.5)
    sec.bottom_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.7)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BASE_SIZE)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(BASE_SIZE * 2))
    rpr.append(szcs)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:bidi"), "th-TH")
    rpr.append(lang)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    # เลขหน้ามุมบนขวา
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run()
    _set_run_font(run)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            run._r.append(it)
    sec.different_first_page_header_footer = True
    return doc


def _set_cell_bg(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def add_table(doc, spec):
    if spec.get("caption"):
        p = doc.add_paragraph()
        _fmt(p, before=6, after=3)
        cap = spec["caption"]
        m = re.match(r"(ตารางที่\s*[\d.]+)\s*(.*)", cap)
        if m:
            add_rich_text(p, m.group(1) + " ", bold=True)
            add_rich_text(p, m.group(2))
        else:
            add_rich_text(p, cap, bold=True)
    headers = spec["headers"]
    rows = spec.get("rows", [])
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = spec.get("widths")
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(p, h, size=spec.get("size", 14), bold=True)
        _set_cell_bg(c, "D9E2F3")
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.text = ""
            p = c.paragraphs[0]
            val = "" if val is None else str(val)
            bold = val.startswith("!")
            if bold:
                val = val[1:]
            align = spec.get("align", ["left"] + ["center"] * (len(headers) - 1))[ci]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
            add_rich_text(p, val, size=spec.get("size", 14), bold=bold)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    if spec.get("note"):
        p = doc.add_paragraph()
        _fmt(p, before=3)
        add_rich_text(p, spec["note"], size=14)
    p = doc.add_paragraph()
    _fmt(p, after=3)
    return t


def _cell_text(cell, text, size=14, bold=False, center=True):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich_text(p, text, size=size, bold=bold)
    return p


def add_gantt(doc, spec):
    """ตารางขั้นตอนการดำเนินงาน: หัวตารางเป็นเดือน (เดือนละ 4 สัปดาห์) แรเงาช่องสัปดาห์ที่ดำเนินการ"""
    months, rows = spec["months"], spec["rows"]
    ncol = len(months) * 4
    t = doc.add_table(rows=2 + len(rows), cols=1 + ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # ช่องสัปดาห์แคบ: ลดระยะขอบในเซลล์เพื่อให้เลขสองหลักอยู่บรรทัดเดียว
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), "40")
        m.set(qn("w:type"), "dxa")
        margins.append(m)
    t._tbl.tblPr.append(margins)
    first = 1.35
    week = (TEXT_WIDTH.inches - first) / ncol
    for row in t.rows:
        row.cells[0].width = Inches(first)
        for c in row.cells[1:]:
            c.width = Inches(week)
    for ci in range(1, ncol + 1):  # เลขสัปดาห์ (ใส่ก่อนผสานเซลล์หัวตาราง)
        _cell_text(t.cell(1, ci), str(ci) if ci <= spec["weeks"] else "", size=10, bold=True)
    for mi, name in enumerate(months):
        merged = t.cell(0, 1 + mi * 4).merge(t.cell(0, 4 + mi * 4))
        _cell_text(merged, name, size=13, bold=True)
    _cell_text(t.cell(0, 0).merge(t.cell(1, 0)), spec.get("label", "กิจกรรม"), size=13, bold=True)
    for ri, (name, start, end) in enumerate(rows, start=2):
        _cell_text(t.cell(ri, 0), name, size=13, center=False)
        for ci in range(1, ncol + 1):
            cell = t.cell(ri, ci)
            cell.text = ""
            if start <= ci <= end:
                _set_cell_bg(cell, "BFBFBF")
    # ให้ทั้งตารางอยู่หน้าเดียวกัน: ทุกย่อหน้าของทุกแถว (ยกเว้นแถวสุดท้าย) ต้องผูกกับแถวถัดไป
    # ไล่จาก XML โดยตรง เพราะเซลล์ที่ถูกผสานแนวตั้งไม่ปรากฏใน row.cells
    for tr in t._tbl.findall(qn("w:tr"))[:-1]:
        for para in tr.iter(qn("w:p")):
            para.get_or_add_pPr().append(OxmlElement("w:keepNext"))
    p = doc.add_paragraph()
    _fmt(p, after=3)
    return t


def add_signatures(doc, spec):
    """ช่องลงนามคู่ซ้าย-ขวา: เส้นประ (ชื่อ) ตำแหน่ง"""
    dots = "." * 60
    t = doc.add_table(rows=len(spec["rows"]), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, pair in enumerate(spec["rows"]):
        for ci, person in enumerate(pair):
            cell = t.cell(ri, ci)
            cell.width = Inches(TEXT_WIDTH.inches / 2)
            if not person.get("role"):
                cell.text = ""
                continue
            p = _cell_text(cell, dots)
            _fmt(p, before=14)
            name = person.get("name") or "." * 50
            add_rich_text(cell.add_paragraph(), f"({name})")
            add_rich_text(cell.add_paragraph(), person["role"])
            for q in cell.paragraphs[1:]:
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t


def add_fill_lines(doc, n, indent=0.0):
    for _ in range(n):
        p = doc.add_paragraph()
        _fmt(p, line=1.15)
        pf = p.paragraph_format
        pf.left_indent = Inches(indent)
        pf.tab_stops.add_tab_stop(TEXT_WIDTH - Inches(0.05), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run("\t")
        _set_run_font(r)


def render(doc, blocks):
    for b in blocks:
        kind = b[0]
        if kind == "chapter":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(p, b[1], size=20, bold=True)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _fmt(p, after=18)
            add_rich_text(p, b[2], size=20, bold=True)
        elif kind == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            size = b[2] if len(b) > 2 else 18
            _fmt(p, after=6)
            add_rich_text(p, b[1], size=size, bold=True)
        elif kind == "center":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(p, b[1], size=b[2] if len(b) > 2 else BASE_SIZE)
        elif kind == "h1":
            p = doc.add_paragraph()
            _fmt(p, before=12, after=3)
            p.paragraph_format.keep_with_next = True
            add_rich_text(p, b[1], size=18, bold=True)
        elif kind == "h2":
            p = doc.add_paragraph()
            _fmt(p, before=6, after=0)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.keep_with_next = True
            add_rich_text(p, b[1], bold=True)
        elif kind == "hd":
            p = doc.add_paragraph()
            _fmt(p, before=12, after=0)
            p.paragraph_format.keep_with_next = True
            add_rich_text(p, b[1], bold=True)
        elif kind in ("p", "pn"):
            p = doc.add_paragraph()
            _thai_justify(p)
            if kind == "p":
                p.paragraph_format.first_line_indent = Inches(0.5)
            add_rich_text(p, b[1])
        elif kind == "item":
            level = b[2] if len(b) > 2 else 1
            p = doc.add_paragraph()
            _thai_justify(p)
            pf = p.paragraph_format
            pf.left_indent = Inches(0.5 + 0.3 * (level - 1) + 0.3)
            pf.first_line_indent = Inches(-0.3)
            add_rich_text(p, b[1])
        elif kind == "table":
            add_table(doc, b[1])
        elif kind == "fill":
            add_fill_lines(doc, b[1], b[2] if len(b) > 2 else 0.0)
        elif kind == "ref":
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.5)
            pf.first_line_indent = Inches(-0.5)
            _fmt(p, after=4)
            add_rich_text(p, b[1])
        elif kind == "gantt":
            add_gantt(doc, b[1])
        elif kind == "sign":
            add_signatures(doc, b[1])
        elif kind == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "blank":
            doc.add_paragraph()
        else:
            raise ValueError(kind)


def build_docx(blocks, out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = new_document()
    render(doc, blocks)
    doc.save(str(out_path))
    return out_path


_word_lock = threading.Lock()


def convert_to_pdf(docx_paths: list[Path]):
    """แปลง .docx -> .pdf ด้วย Microsoft Word (COM) เพื่อให้ฟอนต์ไทยแสดงผลถูกต้อง"""
    with _word_lock:  # Word ทำงานได้ทีละงาน
        _convert_to_pdf(docx_paths)


def _convert_to_pdf(docx_paths: list[Path]):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()  # จำเป็นเมื่อเรียกจาก thread ของเว็บเซิร์ฟเวอร์
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        for d in docx_paths:
            pdf = d.with_suffix(".pdf")
            doc = word.Documents.Open(str(d), ReadOnly=True)
            doc.SaveAs2(str(pdf), FileFormat=17)  # wdFormatPDF
            doc.Close(False)
    finally:
        word.Quit()
