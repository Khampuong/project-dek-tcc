"""รูปแบบบล็อกเนื้อหา: การตรวจรูปร่าง การเรียงเลขหัวข้อ/ตาราง คำสั่งแก้ไขของ AI และการสร้างไฟล์ Word"""
import tempfile
import unittest
from pathlib import Path

from app import editor, store


class SanitizeTest(unittest.TestCase):
    def test_accepts_valid_and_repairs_loose_shapes(self):
        raw = [["h1", "1.1 ก"], {"type": "p", "text": "ย่อหน้า ⟦…⟧"}, "ข้อความล้วน", ["item", "1) ก", "x"],
               ["fill", 99], ["fill", "x"], ["pb"], ["chapter", "บทที่ 1", "บทนำ"], ["zzz", "?"], ["p", "  "], None, 5]
        self.assertEqual(store.sanitize_blocks(raw), [
            ["h1", "1.1 ก"], ["p", "ย่อหน้า ⟦…⟧"], ["p", "ข้อความล้วน"], ["item", "1) ก", 1],
            ["fill", 15], ["fill", 3], ["pb"], ["chapter", "บทที่ 1", "บทนำ"]])

    def test_not_a_list(self):
        self.assertEqual(store.sanitize_blocks({"blocks": []}), [])
        self.assertEqual(store.sanitize_blocks(None), [])

    def test_table_rows_padded_and_truncated(self):
        t = store.sanitize_blocks([["table", {"caption": " ตารางที่ 4.1 ผล ", "headers": ["รายการ", "x̄", "S.D."],
                                              "rows": [["ด้านที่ 1"], ["!รวม", 1, None, "เกิน"], "ไม่ใช่แถว"],
                                              "align": ["left", "bogus", "center"], "size": 99}]])[0][1]
        self.assertEqual(t["rows"], [["ด้านที่ 1", "", ""], ["!รวม", "1", ""]])
        self.assertEqual(t["caption"], "ตารางที่ 4.1 ผล")
        self.assertNotIn("align", t)  # การจัดแนวที่ไม่ถูกต้องถูกทิ้ง ใช้ค่าเริ่มต้น
        self.assertEqual(t["size"], 14)
        self.assertEqual(len(t["widths"]), 3)
        self.assertLessEqual(sum(t["widths"]), store.TEXT_WIDTH_IN + 0.05)

    def test_table_without_headers_dropped(self):
        self.assertEqual(store.sanitize_blocks([["table", {"rows": [["a"]]}], ["table", "x"]]), [])

    def test_oversized_widths_scaled_to_page(self):
        t = store.sanitize_blocks([["table", {"headers": ["a", "b"], "rows": [], "widths": [10, 10]}]])[0][1]
        self.assertAlmostEqual(sum(t["widths"]), store.TEXT_WIDTH_IN, places=1)

    def test_gantt_and_signature_blocks(self):
        out = store.sanitize_blocks([
            ["gantt", {"months": ["ต.ค.", " ", "พ.ย."], "weeks": 99, "rows": [["ก", 2, 1], ["ข", -3, 5], "x", ["", 1, 1]]}],
            ["gantt", {"months": [], "rows": []}],
            ["sign", {"rows": [[{"name": " ก ", "role": "ผู้จัดทำ"}], "x", [{"role": "ครู"}, None, {"role": "เกิน"}]]}],
            ["sign", {"rows": []}]])
        self.assertEqual(out, [
            ["gantt", {"months": ["ต.ค.", "พ.ย."], "weeks": 8, "rows": [["ก", 2, 2], ["ข", 0, 5]], "label": "กิจกรรม"}],
            ["sign", {"rows": [[{"name": "ก", "role": "ผู้จัดทำ"}, {"name": "", "role": ""}],
                               [{"name": "", "role": "ครู"}, {"name": "", "role": ""}]]}]])

    def test_control_characters_stripped(self):
        self.assertEqual(store.sanitize_blocks([["p", "สูตร \x0crac{a}{b}\x00"]]), [["p", "สูตร rac{a}{b}"]])


class RenumberTest(unittest.TestCase):
    def test_headings_made_consecutive(self):
        b = [["h1", "1.1 ก"], ["h1", "1.2 ข"], ["h1", "1.4 ค"], ["h2", "1.5.1 ง"], ["h2", "1.5.2 จ"], ["h1", "1.5 ฉ"], ["h1", "สรุป"]]
        self.assertEqual([x[1] for x in store.renumber_headings(b, 1)],
                         ["1.1 ก", "1.2 ข", "1.3 ค", "1.3.1 ง", "1.3.2 จ", "1.4 ฉ", "สรุป"])

    def test_tables_renumbered_with_text_references(self):
        b = [["p", "ดังตารางที่ 2.3-1 และตารางที่ 2.5-1"],
             ["table", {"caption": "ตารางที่ 2.3-1 ขั้นตอน", "headers": ["a"], "rows": []}],
             ["table", {"caption": "สังเคราะห์งานวิจัย", "headers": ["a"], "rows": []}],
             ["table", {"caption": "ตารางที่ 2.5-1 กรอบ", "headers": ["a"], "rows": []}]]
        out = store.renumber_tables(b, 2)
        self.assertEqual([x[1]["caption"] for x in out[1:]],
                         ["ตารางที่ 2.1 ขั้นตอน", "ตารางที่ 2.2 สังเคราะห์งานวิจัย", "ตารางที่ 2.3 กรอบ"])
        self.assertEqual(out[0][1], "ดังตารางที่ 2.1 และตารางที่ 2.3")


class EditorOpsTest(unittest.TestCase):
    DOC = [["h1", "1.1 ก"], ["p", "หนึ่ง"], ["p", "สอง"], ["p", "สาม"]]

    def test_replace_insert_delete_use_original_positions(self):
        ops = [{"op": "replace", "index": 0, "block": ["h1", "1.1 ใหม่"]},
               {"op": "insert", "after": 1, "blocks": [["p", "แทรก"]]},
               {"op": "delete", "index": 3},
               {"op": "insert", "after": -1, "blocks": [["pn", "ต้นเอกสาร"]]}]
        out, changed = editor.apply_ops(self.DOC, ops)
        self.assertEqual(out, [["pn", "ต้นเอกสาร"], ["h1", "1.1 ใหม่"], ["p", "หนึ่ง"], ["p", "แทรก"], ["p", "สอง"]])
        self.assertEqual(changed, [0, 1, 3])

    def test_invalid_ops_ignored(self):
        ops = [{"op": "replace", "index": 99, "block": ["p", "x"]}, {"op": "replace", "index": 1, "block": ["bogus"]},
               {"op": "explode"}, "not-an-op", {"op": "delete", "index": "abc"}]
        self.assertEqual(editor.apply_ops(self.DOC, ops), (self.DOC, []))


class DocxTest(unittest.TestCase):
    def test_every_block_kind_renders_in_thai_font(self):
        import docx

        from app.docbuilder import build_docx

        blocks = store.sanitize_blocks([
            ["chapter", "บทที่ 1", "บทนำ"], ["title", "ชื่อเรื่อง", 20], ["center", "กึ่งกลาง"], ["h1", "1.1 หัวข้อ"],
            ["hd", "1. หลักการและเหตุผล"],
            ["h2", "1.1.1 ย่อย"], ["p", "ย่อหน้า **หนา** *เอียง* ⟦เติม⟧ ยาวพอให้ตัดคำภาษาไทยทำงาน"], ["pn", "ไม่เยื้อง"],
            ["item", "1) รายการ", 2], ["fill", 2], ["blank"], ["pb"], ["ref", "ก ข. (2567). เรื่อง. *วารสาร, 1*(2), 3–4."],
            ["table", {"caption": "ตารางที่ 1.1 ทดสอบ", "headers": ["รายการ", "x̄"], "rows": [["ด้าน\nที่ 1", ""], ["!รวม", ""]], "note": "หมายเหตุ"}],
            ["gantt", {"months": ["ตุลาคม", "พฤศจิกายน"], "weeks": 6, "rows": [["1. วางแผน", 1, 2], ["2. พัฒนา", 3, 6]]}],
            ["sign", {"rows": [[{"name": "", "role": "ผู้จัดทำโครงงาน"}, {"name": "นางสาวครู ทดสอบ", "role": "ครูประจำวิชา"}]]}]])
        with tempfile.TemporaryDirectory() as d:
            path = build_docx([tuple(b) for b in blocks], Path(d) / "sub" / "t.docx")
            doc = docx.Document(str(path))
        fonts = {r.font.name for p in doc.paragraphs for r in p.runs if r.font.name}
        self.assertEqual(fonts, {"TH SarabunPSK"})
        self.assertEqual(len(doc.tables), 3)
        self.assertEqual(len(doc.tables[0].rows), 3)
        gantt, sign = doc.tables[1], doc.tables[2]
        self.assertEqual((len(gantt.rows), len(gantt.columns)), (4, 9))  # หัวตาราง 2 แถว + 2 กิจกรรม, กิจกรรม + 8 สัปดาห์
        self.assertEqual(gantt.cell(0, 1).text.replace("\u200b", ""), "ตุลาคม")
        self.assertEqual([gantt.cell(1, c).text for c in range(1, 9)], ["1", "2", "3", "4", "5", "6", "", ""])
        shaded = [[bool(gantt.cell(r, c)._tc.xpath(".//w:shd")) for c in range(1, 9)] for r in (2, 3)]
        self.assertEqual(shaded, [[True, True] + [False] * 6, [False, False] + [True] * 4 + [False, False]])
        self.assertIn("(นางสาวครู ทดสอบ)", sign.cell(0, 1).text.replace("\u200b", ""))
        self.assertIn("ผู้จัดทำโครงงาน", sign.cell(0, 0).text.replace("\u200b", ""))
        text = "\n".join(p.text for p in doc.paragraphs).replace("​", "")
        self.assertIn("บทนำ", text)
        self.assertIn("[เติม]", text)  # ช่องเติมข้อมูลแสดงเป็นวงเล็บเหลี่ยมไฮไลต์

    def test_unknown_block_kind_is_an_error(self):
        from app.docbuilder import new_document, render

        with self.assertRaises(ValueError):
            render(new_document(), [("bogus", "x")])


if __name__ == "__main__":
    unittest.main()
